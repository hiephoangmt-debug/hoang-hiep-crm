#!/usr/bin/env python3
"""Điền báo cáo marketing tuần (PPTX mẫu) từ file số liệu xlsx + thư mục ảnh chứng từ.

Ví dụ:
  python fill_report.py --template mau.pptx --data casa.xlsx --week "05/10 - 11/10" \
      --images ./anh --out bao_cao_casa.pptx
Thư mục ảnh: <images>/<facebook|tiktok|zalo|google>/<ads|hoadon>/*.png  (xem README.md)
"""
import argparse
import copy
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.util import Emu

PLATFORMS = ["Facebook", "Tiktok", "Zalo", "Google"]
# từ khoá (đã bỏ dấu, viết thường) -> tên metric
METRIC_KEYS = [
    ("chi phi thuc chay (+", "cost_fee"),
    ("chi phi thuc chay", "cost"),
    ("id tai khoan", "account"),
    ("lead", "lead"),
    ("impressions", "impressions"),
    ("reach", "reach"),
    ("frequency", "frequency"),
    ("luot nhap", "clicks"),
    ("ctr", "ctr"),
    ("cpm", "cpm"),
    ("luot xem", "views"),
]
SUM_FIELDS = ["cost", "cost_fee", "lead", "impressions", "reach", "clicks", "views"]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "").replace("đ", "d").replace("Đ", "D"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def to_num(v):
    """Số kiểu VN ('1.108.866', '1,45') hoặc kiểu US ('97,231', '40,445.52') -> float."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("đ", "").replace("%", "")
    if not re.search(r"\d", s):
        return 0.0
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        parts = s.split(sep)
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and parts[0] != "0"):
            s = s.replace(sep, "")  # dấu phân cách hàng nghìn
        else:
            s = s.replace(sep, ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def fmt_int(x):
    return f"{round(x):,}".replace(",", ".")


def fmt_dec(x, nd=2):
    return f"{x:.{nd}f}".replace(".", ",")


# ---------------------------------------------------------------- đọc số liệu
def load_data(path, sheet=None):
    ws = openpyxl.load_workbook(path, data_only=True)
    ws = ws[sheet] if sheet else ws.worksheets[0]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    hdr_plat = next(i for i, r in enumerate(rows) if sum(str(c).strip() in PLATFORMS for c in r if c) >= 2)
    plat_cols = {str(c).strip(): j for j, c in enumerate(rows[hdr_plat]) if c and str(c).strip() in PLATFORMS}
    order = sorted(plat_cols.items(), key=lambda kv: kv[1])
    spans = {}
    for k, (name, start) in enumerate(order):
        end = order[k + 1][1] if k + 1 < len(order) else start + 11
        spans[name] = (start, min(end, start + 12))
    hdr = rows[hdr_plat + 1]
    colmap = {}  # platform -> metric -> col
    for p, (a, b) in spans.items():
        m = {}
        for j in range(a, b):
            h = norm(hdr[j]) if j < len(hdr) else ""
            for key, name in METRIC_KEYS:
                if h.startswith(key) and name not in m:
                    m[name] = j
                    break
        colmap[p] = m

    weeks, cur = {}, None
    for r in rows[hdr_plat + 2:]:
        label = str(r[0]).strip() if r and r[0] else ""
        if norm(label).startswith("tong"):
            break
        if re.match(r"\d{1,2}/\d{1,2}\s*-\s*\d{1,2}/\d{1,2}", label):
            cur = re.sub(r"\s*-\s*", " - ", label)
            weeks[cur] = {p: [] for p in PLATFORMS}
        if cur is None:
            continue
        for p in PLATFORMS:
            m = colmap.get(p, {})
            if not m:
                continue
            get = lambda k: r[m[k]] if k in m and m[k] < len(r) else None
            rec = {k: to_num(get(k)) for k in SUM_FIELDS}
            rec["account"] = str(get("account") or "").strip()
            rec["frequency"], rec["ctr"], rec["cpm"] = (to_num(get(k)) for k in ("frequency", "ctr", "cpm"))
            if rec["cost"] or rec["cost_fee"] or rec["lead"] or rec["impressions"] or rec["account"]:
                weeks[cur][p].append(rec)
    return weeks


def aggregate(recs):
    t = {k: sum(r[k] for r in recs) for k in SUM_FIELDS}
    t["frequency"] = t["impressions"] / t["reach"] if t["reach"] else 0
    t["ctr"] = t["clicks"] / t["impressions"] if t["impressions"] else 0
    t["cpm"] = t["cost"] / t["impressions"] * 1000 if t["impressions"] else 0
    t["accounts"] = len({r["account"] for r in recs if r["account"]})
    return t


# ------------------------------------------------------------------ điền PPTX
def shapes(slide):
    for sh in slide.shapes:
        if sh.shape_type == 6:  # group
            yield from sh.shapes
        else:
            yield sh


def slide_text(slide):
    return "\n".join(sh.text_frame.text for sh in shapes(slide) if sh.has_text_frame)


def slide_platform(slide):
    """Nền tảng của slide theo tiêu đề (FACEBOOK / TIKTOK ADS / ZALO / GOOGLE ...)."""
    lines = [norm(l) for l in slide_text(slide).split("\n")]
    for p in PLATFORMS:
        if any(re.fullmatch(rf"{p.lower()}( ads)?:?", l) for l in lines):
            return p
    return None


def set_text(shape, text):
    """Thay text giữ nguyên định dạng run đầu tiên."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def set_cell(cell, text):
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text


def table_header(tbl):
    return [norm(c.text) for c in tbl.rows[0].cells]


def col_of(header, *keys):
    for j, h in enumerate(header):
        if any(h.startswith(k) for k in keys):
            return j
    return None


def row_values(header, p, rec, order):
    """Giá trị từng cột của bảng mẫu cho một dòng (rec=None -> dòng trống)."""
    r = rec or {}
    vals = []
    for h in header:
        if h.startswith("nen tang"):
            v = p
        elif h.startswith("id tai khoan"):
            v = r.get("account", "")
        elif h.startswith("chi phi thuc chay ("):
            v = fmt_int(r["cost_fee"]) if r else ""
        elif h.startswith("ngan sach thuc chay"):
            v = fmt_int(r["cost"]) if r else ""
        elif h.startswith("lead"):
            v = fmt_int(r["lead"]) if r else ""
        elif h.startswith("impressions"):
            v = fmt_int(r["impressions"]) if r else ""
        elif h.startswith("reach"):
            v = fmt_int(r["reach"]) if r else ""
        elif h.startswith("frequency"):
            v = fmt_dec(r["frequency"]) if r else ""
        elif h.startswith("luot nhap"):
            v = fmt_int(r["clicks"]) if r else ""
        elif h.startswith("ctr"):
            v = fmt_dec(r["ctr"] * 100) + "%" if r else ""
        elif h.startswith("cpm"):
            v = fmt_int(r["cpm"]) if r else ""
        elif h.startswith("luot xem"):
            v = fmt_int(r["views"]) if r else ""
        elif h in ("deal", "mql ( khach hang quan tam )", "sql ( khach hang tiem nang)") or h.startswith(("deal", "mql", "sql")):
            v = "-" if r else ""
        else:
            v = ""
        vals.append(v)
    return vals


def ensure_rows(tbl, n):
    """Đảm bảo bảng có đủ n dòng dữ liệu (nhân bản dòng cuối)."""
    while len(tbl.rows) - 1 < n:
        tr = copy.deepcopy(tbl._tbl.tr_lst[-1])
        tbl._tbl.append(tr)


def fill_table(tbl, platform_filter, data_by_platform):
    header = table_header(tbl)
    lines = []
    for p in PLATFORMS:
        if platform_filter and p != platform_filter:
            continue
        recs = [r for r in data_by_platform.get(p, [])]
        if not recs:
            continue
        # mỗi tài khoản (ID) một dòng, gộp theo ID
        by_acc = {}
        for r in recs:
            by_acc.setdefault(r["account"], []).append(r)
        for acc, rs in by_acc.items():
            a = aggregate(rs)
            a["account"] = acc
            lines.append((p, a))
    # dòng tổng
    if len(lines) > 1 or not platform_filter:
        allrecs = [r for p in PLATFORMS if not platform_filter or p == platform_filter for r in data_by_platform.get(p, [])]
        if allrecs:
            t = aggregate(allrecs)
            t["account"] = ""
            lines.append(("TỔNG", t))
    if not lines:
        return
    ensure_rows(tbl, len(lines))
    for i, (p, rec) in enumerate(lines, start=1):
        for j, v in enumerate(row_values(header, p, rec, None)):
            set_cell(tbl.cell(i, j), v)
    for i in range(len(lines) + 1, len(tbl.rows)):  # xoá dòng thừa
        for j in range(len(header)):
            set_cell(tbl.cell(i, j), "")


def fill_text(slide, ctx):
    for sh in shapes(slide):
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if re.fullmatch(r"TUẦN\s*\d*", t, re.I) and ctx.get("week_no"):
            set_text(sh, f"TUẦN {ctx['week_no']}")
        elif re.fullmatch(r"\.{3,}\s*\d{1,2}\.\d{4}\s*-\s*\.{3,}\s*\d{1,2}\.\d{4}", t):
            set_text(sh, ctx["week_range_text"])
        elif t.startswith("Thời gian triển khai:") and ctx.get("period"):
            set_text(sh, f"Thời gian triển khai: {ctx['period']}")
        elif t.startswith("REVIEW CHIẾN DỊCH") and ctx.get("campaign"):
            set_text(sh, f"REVIEW CHIẾN DỊCH {ctx['campaign']}")
        else:
            m = re.fullmatch(r"(Facebook|Tiktok|Zalo|Google):", t, re.I)
            if m:
                p = m.group(1).capitalize()
                set_text(sh, f"{p}: {fmt_int(ctx['totals'][p]['cost_fee'])} VNĐ (đã gồm phí nền tảng)")
            elif t.startswith("NHÂN SỰ:") and ctx.get("staff"):
                set_text(sh, f"NHÂN SỰ: {ctx['staff']}")


def add_images(prs, slide, platform, kind, folder):
    files = sorted(p for p in Path(folder).glob("*") if p.suffix.lower() in (".png", ".jpg", ".jpeg"))
    if not files:
        return 0
    W, H = prs.slide_width, prs.slide_height
    top, margin = int(H * 0.30), int(W * 0.03)
    n = len(files)
    cols = 1 if n == 1 else 2
    rows_ = -(-n // cols)
    cw = (W - margin * (cols + 1)) // cols
    ch = (int(H * 0.92) - top - margin * (rows_ - 1)) // rows_
    for k, f in enumerate(files):
        c, r = k % cols, k // cols
        pic = slide.shapes.add_picture(str(f), 0, 0)
        ratio = min(cw / pic.width, ch / pic.height)
        pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        pic.left = margin + c * (cw + margin) + (cw - pic.width) // 2
        pic.top = top + r * (ch + margin) + (ch - pic.height) // 2
    return n


def slide_kind(slide):
    t = norm(slide_text(slide))
    if "hoa don" in t:
        return "hoadon"
    if "demo:" in t or "capture" in t:
        return "ads"
    return None


def build(args):
    weeks = load_data(args.data, args.sheet)
    active = [w for w, d in weeks.items() if any(d[p] for p in PLATFORMS)]
    if not active:
        sys.exit("Không có dữ liệu thực chạy trong file số liệu.")
    week = args.week or active[-1]
    if week not in weeks:
        sys.exit(f"Không tìm thấy tuần '{week}'. Các tuần có: {', '.join(weeks)}")
    week_data = weeks[week]
    total_data = {p: [r for w in weeks.values() for r in w[p]] for p in PLATFORMS}
    totals = {p: aggregate(total_data[p]) for p in PLATFORMS}
    weekrange = week.replace(" - ", " - ")
    d1, d2 = [x.strip() for x in week.split("-")]
    year = args.year
    ctx = dict(
        totals=totals, campaign=args.campaign, period=args.period, staff=args.staff,
        week_no=args.week_no or (list(weeks).index(week) + 1 - (list(weeks).index(active[0]))),
        week_range_text=f"{d1.replace('/', '.')}.{year} - {d2.replace('/', '.')}.{year}",
    )
    prs = Presentation(args.template)
    imgs = Path(args.images) if args.images else None
    n_img = 0
    for slide in prs.slides:
        plat = slide_platform(slide)
        fill_text(slide, ctx)
        txt = norm(slide_text(slide))
        for sh in shapes(slide):
            if not sh.has_table:
                continue
            tbl = sh.table
            header = table_header(tbl)
            weekly = col_of(header, "deal") is not None and col_of(header, "deal") < col_of(header, "lead", "lead") if "lead" in " ".join(header) else False
            fill_table(tbl, plat, week_data if weekly else total_data)
        kind = slide_kind(slide)
        if imgs and plat and kind:
            n_img += add_images(prs, slide, plat, kind, imgs / plat.lower() / kind)
    prs.save(args.out)
    print(f"Đã tạo {args.out} | tuần: {week} | ảnh chèn: {n_img}")
    for p in PLATFORMS:
        t = totals[p]
        print(f"  {p:9s} chi phí={fmt_int(t['cost'])} (+phí {fmt_int(t['cost_fee'])}) lead={fmt_int(t['lead'])}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", required=True, help="File PPTX mẫu")
    ap.add_argument("--data", required=True, help="File xlsx số liệu (tải từ Google Sheet: File > Tải xuống > .xlsx)")
    ap.add_argument("--sheet", help="Tên sheet (mặc định: sheet đầu)")
    ap.add_argument("--week", help="Tuần báo cáo, ví dụ '05/10 - 11/10' (mặc định: tuần cuối có dữ liệu)")
    ap.add_argument("--week-no", type=int, help="Số thứ tự tuần ghi trên slide (mặc định: tự tính)")
    ap.add_argument("--year", default="2026")
    ap.add_argument("--campaign", help="Tên chiến dịch, ví dụ 'CASAMIA BALANCA HOI AN-TATI INVEST ĐỢT 1'")
    ap.add_argument("--period", help="Thời gian triển khai, ví dụ '15/09 - 27/09/2026'")
    ap.add_argument("--staff", help="Nhân sự phụ trách")
    ap.add_argument("--images", help="Thư mục ảnh chứng từ: <dir>/<nền tảng>/<ads|hoadon>/*.png")
    ap.add_argument("--out", default="bao_cao.pptx")
    build(ap.parse_args())


if __name__ == "__main__":
    main()
